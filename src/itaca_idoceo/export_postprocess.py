from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from .core import safe_filename_component
from .names import normalize_person_name


_COLLISION_SUFFIX_RE = re.compile(r"(__\d+)$")


def normalize_xlsx_names(path: Path) -> None:
    """Normaliza sólo las columnas Apellidos/Nombre de un XLSX ya generado."""
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        headers = {
            str(cell.value).strip(): cell.column
            for cell in sheet[1]
            if cell.value is not None
        }
        surname_col = headers.get("Apellidos")
        given_col = headers.get("Nombre")
        if surname_col is None or given_col is None:
            continue

        for row in range(2, sheet.max_row + 1):
            surname_cell = sheet.cell(row=row, column=surname_col)
            given_cell = sheet.cell(row=row, column=given_col)
            surnames = str(surname_cell.value or "")
            given = str(given_cell.value or "")
            normalized_surnames, normalized_given = normalize_person_name(
                surnames,
                given,
            )
            surname_cell.value = normalized_surnames
            given_cell.value = normalized_given

    workbook.save(path)


def _normalized_photo_stem(stem: str) -> str:
    suffix = ""
    match = _COLLISION_SUFFIX_RE.search(stem)
    if match:
        suffix = match.group(1)
        stem = stem[: -len(suffix)]

    if "," not in stem:
        return safe_filename_component(stem) + suffix

    surnames, given = stem.split(",", 1)
    normalized_surnames, normalized_given = normalize_person_name(
        surnames.strip(),
        given.strip(),
    )
    return safe_filename_component(
        f"{normalized_surnames}, {normalized_given}"
    ) + suffix


def normalize_name_photo_filenames(folder: Path | None) -> None:
    """Alinea los nombres de foto por nombre con la capitalización del XLSX."""
    if folder is None or not folder.is_dir():
        return

    files = sorted(
        (path for path in folder.iterdir() if path.is_file()),
        key=lambda path: path.name.casefold(),
    )
    if not files:
        return

    # Dos fases para que los cambios sólo de mayúsculas funcionen también en
    # sistemas de ficheros no sensibles a mayúsculas/minúsculas.
    temporary: list[tuple[Path, Path]] = []
    for index, source in enumerate(files, start=1):
        target_name = _normalized_photo_stem(source.stem) + source.suffix.lower()
        temp = folder / f".__itaca_idoceo_tmp_{index}{source.suffix.lower()}"
        source.rename(temp)
        temporary.append((temp, folder / target_name))

    used: set[str] = set()
    for temp, desired in temporary:
        candidate = desired
        counter = 2
        while candidate.name.casefold() in used or candidate.exists():
            candidate = desired.with_name(
                f"{desired.stem}__{counter}{desired.suffix}"
            )
            counter += 1
        temp.rename(candidate)
        used.add(candidate.name.casefold())
