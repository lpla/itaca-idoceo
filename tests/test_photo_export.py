from __future__ import annotations

import struct
import zlib
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

import itaca_idoceo.photo_export as photo_export
from itaca_idoceo.core import ClassResult, PageMetadata, Student
from itaca_idoceo.photo_match import PhotoRosterReferenceMatch, PhotoStudentLink
from itaca_idoceo.photo_reference_pool import ReferencePoolMatch
from itaca_idoceo.photo_roster import (
    PhotoRosterPageSummary,
    PhotoRosterResult,
    PhotoRosterStudent,
)


def _png_rgb(width: int = 2, height: int = 2) -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        payload = kind + data
        return (
            struct.pack(">I", len(data))
            + payload
            + struct.pack(">I", zlib.crc32(payload) & 0xFFFFFFFF)
        )

    signature = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    row = b"\x00" + (b"\xff\xff\xff" * width)
    pixels = row * height
    return (
        signature
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(pixels))
        + chunk(b"IEND", b"")
    )


def _make_source_pdf(path) -> int:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    xref = page.insert_image(
        pymupdf.Rect(25, 178, 102, 273),
        stream=_png_rgb(),
    )
    document.save(path)
    document.close()
    return xref


def _roster_result(pdf_path, xref: int, *, second_has_photo: bool = False) -> PhotoRosterResult:
    return PhotoRosterResult(
        source=pdf_path,
        page_count=1,
        students=[
            PhotoRosterStudent(
                page=1,
                ordinal=1,
                row=1,
                column=1,
                full_name="GARCIA LOPEZ, ANA",
                surnames="GARCIA LOPEZ",
                given_names="ANA",
                xref=xref,
                image_x0=25,
                image_y0=178,
                image_x1=102,
                image_y1=273,
                name_lines=1,
            ),
            PhotoRosterStudent(
                page=1,
                ordinal=2,
                row=1,
                column=2,
                full_name="PEREZ, BEA",
                surnames="PEREZ",
                given_names="BEA",
                xref=xref if second_has_photo else 0,
                image_x0=119,
                image_y0=178,
                image_x1=196,
                image_y1=273,
                name_lines=1,
            ),
        ],
        pages=[
            PhotoRosterPageSummary(
                1,
                2 if second_has_photo else 1,
                2,
                1,
                2,
                0,
                0 if second_has_photo else 1,
            )
        ],
        issues=[],
        group_raw="3ESO A - 3ESO A",
        group_code="3ESO A",
        tutor="Doe, Jane",
    )


def _reference_student(
    ordinal: int,
    nia: str,
    surnames: str,
    given_names: str,
    repetix: str,
    materia: str,
) -> Student:
    return Student(
        page=1,
        block=ordinal,
        ordinal=ordinal,
        nia=nia,
        full_name=f"{surnames}, {given_names}",
        surnames=surnames,
        given_names=given_names,
        row_x=0.0,
        visual_y=float(ordinal),
        repetix=repetix,
        materia=materia,
    )


def _pool(roster, match, *, unique_students: int = 2) -> ReferencePoolMatch:
    return ReferencePoolMatch(
        photo_result=roster,
        match=match,
        input_files=2,
        usable_files=2,
        skipped_files=0,
        valid_classes=2,
        unique_students=unique_students,
        same_group_students=unique_students,
    )


def test_export_photo_roster_keeps_student_without_available_photo(tmp_path, monkeypatch):
    pdf_path = tmp_path / "grupo_materia.pdf"
    output_dir = tmp_path / "salida"
    xref = _make_source_pdf(pdf_path)
    roster = _roster_result(pdf_path, xref)
    monkeypatch.setattr(photo_export, "detect_photo_roster", lambda _: roster)

    result = photo_export.export_photo_roster(pdf_path, output_dir)

    assert result.students == 2
    assert result.photos == 1
    assert result.missing_photos == 1
    assert result.match_field == "Nombre"
    assert result.automatic_matches == 1
    assert result.manual_matches == 0
    assert result.photos_dir.name == "fotos_por_nombre"
    assert result.name_photos_dir == result.photos_dir

    workbook = load_workbook(result.xlsx_path)
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == ["Apellidos", "Nombre"]
    assert [sheet["A2"].value, sheet["B2"].value] == ["GARCIA LOPEZ", "ANA"]
    assert [sheet["A3"].value, sheet["B3"].value] == ["PEREZ", "BEA"]

    photos = sorted(path.name for path in result.photos_dir.glob("*.png"))
    assert photos == ["GARCIA LOPEZ, ANA.png"]
    assert (result.photos_dir / photos[0]).stat().st_size > 0

    guide = result.guide_path.read_text(encoding="utf-8")
    assert "FOTOS POR NOMBRE" in guide
    assert "asociación por nombre puede fallar" in guide
    assert "sin fotografía disponible en el PDF: 1" in guide


def test_export_with_references_uses_nia_for_xlsx_and_photo_names(tmp_path, monkeypatch):
    pdf_path = tmp_path / "grupo_materia.pdf"
    references = [tmp_path / "a.pdf", tmp_path / "b.pdf"]
    for path in references:
        path.write_bytes(b"dummy")
    output_dir = tmp_path / "salida"
    xref = _make_source_pdf(pdf_path)
    roster = _roster_result(pdf_path, xref)
    monkeypatch.setattr(photo_export, "detect_photo_roster", lambda _: roster)

    first = _reference_student(1, "12345678", "GARCIA LOPEZ", "ANA", "R", "MUSICA")
    second = _reference_student(2, "87654321", "PEREZ", "BEA", "", "MUSICA")
    reference_class = ClassResult(
        source=Path("listado_tabular.pdf"),
        metadata=PageMetadata(group_code="3ESO A"),
        students=[first, second],
        issues=[],
    )
    match = PhotoRosterReferenceMatch(
        photo_result=roster,
        reference_class=reference_class,
        links=[
            PhotoStudentLink(roster.students[0], first, "exact"),
            PhotoStudentLink(roster.students[1], second, "exact"),
        ],
        unmatched_photos=0,
        ambiguous_photos=0,
        issues=[],
    )
    monkeypatch.setattr(
        photo_export,
        "match_photo_result_to_references",
        lambda _roster, _references: _pool(roster, match),
    )

    result = photo_export.export_photo_roster(
        pdf_path,
        output_dir,
        reference_pdfs=references,
        include_repetix=True,
        include_materia=True,
    )

    assert result.students == 2
    assert result.matched_students == 2
    assert result.unmatched_students == 0
    assert result.photos == 1
    assert result.missing_photos == 1
    assert result.match_field == "NIA"
    assert result.automatic_matches == 1
    assert result.manual_matches == 0
    assert result.photos_dir.name == "fotos_por_nia"
    assert result.id_photos_dir == result.photos_dir

    workbook = load_workbook(result.xlsx_path)
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == [
        "Apellidos",
        "Nombre",
        "NIA",
        "REPETIX",
        "MATÈRIA",
    ]
    assert [sheet["C2"].value, sheet["D2"].value, sheet["E2"].value] == [
        "12345678",
        "R",
        "MUSICA",
    ]
    assert sheet["C3"].value == "87654321"

    photos = sorted(path.name for path in result.photos_dir.glob("*.png"))
    assert photos == ["12345678.png"]

    guide = result.guide_path.read_text(encoding="utf-8")
    assert "NIA al campo ID / Student ID" in guide
    assert "FOTOS POR NIA" in guide


def test_export_with_references_allows_one_truly_unmatched_student(tmp_path, monkeypatch):
    pdf_path = tmp_path / "grupo_materia.pdf"
    reference_pdf = tmp_path / "referencia.pdf"
    reference_pdf.write_bytes(b"dummy")
    output_dir = tmp_path / "salida"
    xref = _make_source_pdf(pdf_path)
    roster = _roster_result(pdf_path, xref)
    monkeypatch.setattr(photo_export, "detect_photo_roster", lambda _: roster)

    first = _reference_student(1, "12345678", "GARCIA LOPEZ", "ANA", "R", "MUSICA")
    reference_class = ClassResult(
        source=reference_pdf,
        metadata=PageMetadata(group_code="3ESO A"),
        students=[first],
        issues=[],
    )
    match = PhotoRosterReferenceMatch(
        photo_result=roster,
        reference_class=reference_class,
        links=[PhotoStudentLink(roster.students[0], first, "exact")],
        unmatched_photos=1,
        ambiguous_photos=0,
        issues=["No se ha podido cruzar todo el alumnado del listado con fotos de forma inequívoca"],
    )
    monkeypatch.setattr(
        photo_export,
        "match_photo_result_to_references",
        lambda _roster, _references: _pool(roster, match, unique_students=1),
    )

    result = photo_export.export_photo_roster(
        pdf_path,
        output_dir,
        reference_pdfs=[reference_pdf],
        include_repetix=True,
    )

    assert result.matched_students == 1
    assert result.unmatched_students == 1
    assert result.automatic_matches == 1
    assert result.manual_matches == 0
    assert result.manual_photos_dir is None

    workbook = load_workbook(result.xlsx_path)
    sheet = workbook.active
    assert sheet["C2"].value == "12345678"
    assert sheet["C3"].value is None
    assert sheet["D3"].value is None

    guide = result.guide_path.read_text(encoding="utf-8")
    assert "Alumnos sin coincidencia inequívoca en las referencias: 1" in guide
    assert "NIA vacío" in guide


def test_export_unmatched_student_with_photo_goes_to_name_folder(tmp_path, monkeypatch):
    pdf_path = tmp_path / "grupo_materia.pdf"
    reference_pdf = tmp_path / "referencia.pdf"
    reference_pdf.write_bytes(b"dummy")
    output_dir = tmp_path / "salida"
    xref = _make_source_pdf(pdf_path)
    roster = _roster_result(pdf_path, xref, second_has_photo=True)
    monkeypatch.setattr(photo_export, "detect_photo_roster", lambda _: roster)

    first = _reference_student(1, "12345678", "GARCIA LOPEZ", "ANA", "", "MUSICA")
    reference_class = ClassResult(
        source=reference_pdf,
        metadata=PageMetadata(group_code="3ESO A"),
        students=[first],
        issues=[],
    )
    match = PhotoRosterReferenceMatch(
        photo_result=roster,
        reference_class=reference_class,
        links=[PhotoStudentLink(roster.students[0], first, "exact")],
        unmatched_photos=1,
        ambiguous_photos=0,
        issues=["incomplete"],
    )
    monkeypatch.setattr(
        photo_export,
        "match_photo_result_to_references",
        lambda _roster, _references: _pool(roster, match, unique_students=1),
    )

    result = photo_export.export_photo_roster(
        pdf_path,
        output_dir,
        reference_pdfs=[reference_pdf],
    )

    assert result.photos == 2
    assert result.automatic_matches == 1
    assert result.manual_matches == 1
    assert result.manual_photos_dir is not None
    assert result.manual_photos_dir.name == "fotos_por_nombre"
    assert result.name_photos_dir == result.manual_photos_dir
    assert [path.name for path in result.photos_dir.glob("*.png")] == ["12345678.png"]
    assert [path.name for path in result.manual_photos_dir.glob("*.png")] == ["PEREZ, BEA.png"]

    guide = result.guide_path.read_text(encoding="utf-8")
    assert "FOTOS POR NOMBRE — COMPROBAR" in guide
