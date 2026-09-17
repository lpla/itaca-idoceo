from pathlib import Path

from openpyxl import Workbook, load_workbook

from itaca_idoceo.export_postprocess import (
    normalize_name_photo_filenames,
    normalize_xlsx_names,
)
from itaca_idoceo.names import normalize_name_field, normalize_person_name


def test_normalize_name_field_handles_iberian_particles_and_accents():
    assert normalize_name_field("GARCÍA DE LA FUENTE") == "García de la Fuente"
    assert normalize_name_field("MARÍA DE LOS ÁNGELES") == "María de los Ángeles"
    assert normalize_name_field("DA SILVA") == "da Silva"


def test_normalize_name_field_handles_joiners_and_preserves_existing_case():
    assert normalize_name_field("O'NEILL") == "O'Neill"
    assert normalize_name_field("MARÍA-JOSÉ") == "María-José"
    assert normalize_name_field("MCDONALD") == "McDonald"
    assert normalize_name_field("McDonald") == "McDonald"
    assert normalize_name_field("III") == "III"


def test_normalize_person_name_keeps_fields_separate():
    assert normalize_person_name("DE LA TORRE", "ANA MARÍA") == (
        "de la Torre",
        "Ana María",
    )


def test_normalize_xlsx_names_only_changes_name_columns(tmp_path):
    path = tmp_path / "alumnado.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Apellidos", "Nombre", "NIA", "MATÈRIA"])
    sheet.append(["GARCÍA DE LA FUENTE", "MARÍA-JOSÉ", "12345", "MÚSICA"])
    workbook.save(path)

    normalize_xlsx_names(path)

    result = load_workbook(path).active
    assert result["A2"].value == "García de la Fuente"
    assert result["B2"].value == "María-José"
    assert result["C2"].value == "12345"
    assert result["D2"].value == "MÚSICA"


def test_normalize_name_photo_filenames_preserves_duplicate_suffix(tmp_path):
    folder = tmp_path / "fotos_por_nombre"
    folder.mkdir()
    (folder / "GARCÍA DE LA FUENTE, ANA MARÍA.png").write_bytes(b"a")
    (folder / "O'NEILL, MARÍA-JOSÉ__2.png").write_bytes(b"b")

    normalize_name_photo_filenames(folder)

    names = sorted(path.name for path in folder.iterdir())
    assert names == [
        "García de la Fuente, Ana María.png",
        "O'Neill, María-José__2.png",
    ]
